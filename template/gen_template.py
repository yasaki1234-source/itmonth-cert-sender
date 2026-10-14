"""產生「大專碩博AI×學習產業見習證書」A4 橫式範本 (pptx)。
用法: python3 gen_template.py <輸出路徑> [preview]
preview 模式改用本機有的 Noto CJK 字型名稱，只供本機預覽用。
"""
import sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

out = sys.argv[1]
preview = len(sys.argv) > 2 and sys.argv[2] == "preview"
SERIF = "Noto Serif CJK TC" if preview else "Noto Serif TC"
SANS = "Noto Sans CJK TC" if preview else "Noto Sans TC"

NAVY = RGBColor(0x1F, 0x3A, 0x66)
GOLD = RGBColor(0xCF, 0xA6, 0x50)
GOLD_LIGHT = RGBColor(0xE6, 0xC8, 0x78)
GOLD_DEEP = RGBColor(0xB0, 0x86, 0x38)
RED = RGBColor(0x96, 0x1B, 0x20)
INK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BOX_LINE = RGBColor(0xD8, 0xB9, 0x87)

prs = Presentation()
prs.slide_width = Cm(29.7)
prs.slide_height = Cm(21.0)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版面
W = 29.7


def rect(shape_type, x, y, w, h, fill, line=None, line_w=None):
    s = slide.shapes.add_shape(shape_type, Cm(x), Cm(y), Cm(w), Cm(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w or 0.75)
    return s


def set_run(run, text, font, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def textbox(x, y, w, h, lines, font, size, color, bold=False,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, name=None):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        set_run(p.add_run(), line, font, size, color, bold)
    return tb


# ── 上方深藍標題區與金線 ─────────────────────────────
rect(MSO_SHAPE.RECTANGLE, 0, 0, W, 7.6, NAVY)
rect(MSO_SHAPE.RECTANGLE, 0, 7.6, W, 0.36, GOLD)

# Logo 位置（先放文字，請在 Google 簡報換成正式 Logo 圖檔）
textbox(0.9, 0.55, 12, 1.0, ["資訊月 × 臺灣教育科技展"], SANS, 15, WHITE, bold=True,
        align=PP_ALIGN.LEFT, name="LOGO_請換成圖檔")

# 主標題
textbox(3.0, 2.0, W - 6.0, 4.9, ["大專碩博AI x學習產業", "見習證書"], SERIF, 46, WHITE,
        bold=True, spacing=1.12, name="主標題")

# ── 獎章（緞帶在後、章在前）──────────────────────────
cx, cy = 25.45, 7.5
rect(MSO_SHAPE.RECTANGLE, cx - 1.55, cy + 0.6, 3.1, 3.3, RED)
notch = rect(MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.75, cy + 3.36, 1.5, 0.56, WHITE)
rect(MSO_SHAPE.STAR_24_POINT, cx - 2.2, cy - 2.2, 4.4, 4.4, GOLD)
rect(MSO_SHAPE.OVAL, cx - 1.88, cy - 1.88, 3.76, 3.76, GOLD_LIGHT)
rect(MSO_SHAPE.OVAL, cx - 1.66, cy - 1.66, 3.32, 3.32, GOLD_DEEP)

# ── 內文 ───────────────────────────────────────────
textbox(W / 2 - 4, 9.05, 8, 0.9, ["茲證明"], SANS, 13, GRAY)
textbox(6.25, 10.18, 3, 0.5, ["姓名"], SANS, 8, GOLD_DEEP, align=PP_ALIGN.LEFT)
box = rect(MSO_SHAPE.ROUNDED_RECTANGLE, 6.2, 10.72, 17.3, 2.7, None, BOX_LINE, 1)
box.adjustments[0] = 0.06
textbox(6.4, 10.72, 16.9, 2.7, ["{{姓名}}"], SERIF, 50, INK, name="姓名")

textbox(3.0, 13.85, W - 6.0, 1.8,
        ["參與「{{年度}}資訊月暨臺灣教育科技展」，取得",
         "AI × 學習產業見習時數{{時數}}小時，特頒此證。"],
        SANS, 13, GRAY, spacing=1.25, name="內文")

# ── 底部：日期／證書編號／主辦單位 ─────────────────────
textbox(6.8, 17.3, 5.1, 1.0, ["{{發證年月}}"], SANS, 14, INK, name="發證年月")
rect(MSO_SHAPE.RECTANGLE, 6.85, 18.42, 5.0, 0.03, GRAY)

textbox(W / 2 - 2.8, 17.05, 5.6, 0.6, ["證書編號"], SANS, 8, GOLD_DEEP)
textbox(W / 2 - 2.8, 17.6, 5.6, 0.8, ["{{證書編號}}"], SANS, 11, INK, name="證書編號")

textbox(17.0, 16.85, 6.5, 1.5, ["台北市電腦商業同業公會", "資訊月活動委員會"],
        SANS, 13, INK, bold=True, spacing=1.15, name="主辦單位")
rect(MSO_SHAPE.RECTANGLE, 17.75, 18.42, 5.0, 0.03, GRAY)

prs.save(out)
print("saved", out)
